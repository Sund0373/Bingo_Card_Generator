"""
Creates the input Excel template that users fill in with their 75 bingo terms.

Run once (or whenever you need a fresh blank template):
    python make_template.py            # refuses to clobber an existing template
    python make_template.py --force    # overwrite it anyway
"""

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

TEMPLATE_PATH = Path(__file__).parent / "templates" / "terms_template.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)


def build_template() -> None:
    wb = Workbook()

    # --- Sheet 1: Column-based layout (B/I/N/G/O, 15 terms each) ---
    ws1 = wb.active
    ws1.title = "By Column"
    headers = ["B", "I", "N", "G", "O"]
    for col, letter in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=col, value=letter)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        ws1.column_dimensions[cell.column_letter].width = 22

    for row in range(2, 17):  # rows 2-16 = 15 terms per column
        for col in range(1, 6):
            ws1.cell(row=row, column=col)

    # --- Sheet 2: Flat layout (single column, 75 terms, unassigned) ---
    ws2 = wb.create_sheet("Flat List")
    cell = ws2.cell(row=1, column=1, value="Term")
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    ws2.column_dimensions["A"].width = 30
    for row in range(2, 77):  # rows 2-76 = 75 terms
        ws2.cell(row=row, column=1)

    # --- Instructions sheet ---
    ws3 = wb.create_sheet("Instructions")
    ws3.column_dimensions["A"].width = 100
    lines = [
        "BINGO TERM TEMPLATE — INSTRUCTIONS",
        "",
        "Fill in EITHER sheet below, or both. You always choose a draw style (Column or "
        "Random in the UI, or --mode on the command line), and that choice decides which "
        "sheet is read. Nothing is ever guessed from which sheet you filled in.",
        "",
        "Terms are converted to ALL CAPS automatically, so type them in any case you like.",
        "",
        "All terms must be UNIQUE once upper-cased ('PTO' and 'pto' count as the same term). "
        "Duplicates are rejected with an error naming them, since a repeated term could "
        "otherwise land twice on one card.",
        "",
        "1) 'By Column' sheet  ->  read by the COLUMN draw style",
        "   - Fill in exactly 15 terms under each of B / I / N / G / O.",
        "   - Each card draws 5 terms from that column's own 15, without replacement "
        "(4 for the centre column when a free space is used).",
        "   - A term only ever appears in the column you assigned it to.",
        "",
        "2) 'Flat List' sheet  ->  read by the RANDOM draw style",
        "   - Fill in 75 terms in a single column, in any order.",
        "   - Each card draws 25 terms (24 with a free space) from all 75 at once, without "
        "replacement. Terms are not tied to any column and can appear anywhere, and the "
        "placement is redrawn independently for every card.",
        "",
        "Once filled in, save this file and upload it in the UI, or pass its path to "
        "generate_bingo.py with --terms and --mode column | random.",
    ]
    for i, line in enumerate(lines, start=1):
        c = ws3.cell(row=i, column=1, value=line)
        if i == 1:
            c.font = Font(bold=True, size=13)

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TEMPLATE_PATH)
    print(f"Template created at: {TEMPLATE_PATH}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Create the blank terms template.")
    parser.add_argument(
        "--force", action="store_true", help="Overwrite an existing template file"
    )
    args = parser.parse_args()

    # The template is where people type their terms, so overwriting it destroys work.
    if TEMPLATE_PATH.exists() and not args.force:
        sys.exit(
            f"{TEMPLATE_PATH} already exists and may contain your terms.\n"
            "Refusing to overwrite it. Either back it up and rerun with --force,\n"
            "or delete the file first."
        )

    build_template()


if __name__ == "__main__":
    main()
