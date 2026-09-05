"""
Renders generated bingo cards into a Word document, using the user's
'Bingo Layout.docx' as the formatting source.

The template is not rebuilt from scratch — its first 5x6 table is cloned for every
card, so header shading, fonts, borders, row heights and the FREE SPACE cell keep
their exact original formatting. The template's section is already landscape with
two text columns, so appending a column break after each card yields two cards per
page automatically.
"""

import copy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

TERM_FONT = "Arial"
TERM_SIZE_PT = 14
TERM_COLOR = RGBColor(0x00, 0x00, 0x00)

HEADER_ROW = 0          # table row 0 holds the C/H/A/S/E letters
GRID_ROW_OFFSET = 1     # table row 1 is grid row 0
FREE_ROW, FREE_COL = 2, 2  # grid coordinates of the centre square


def _clear_cell(cell) -> None:
    """Drops every paragraph but the first, and every run in it."""
    for para in cell.paragraphs[1:]:
        para._element.getparent().remove(para._element)
    para = cell.paragraphs[0]
    for run in para.runs:
        run._element.getparent().remove(run._element)


def _remove_shading(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for shd in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(shd)


def _set_term(cell, text: str) -> None:
    """Writes a term into a grid cell as 14pt Arial black, centred."""
    _clear_cell(cell)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.bold = False
    run.font.name = TERM_FONT
    run.font.size = Pt(TERM_SIZE_PT)
    run.font.color.rgb = TERM_COLOR
    # python-docx sets ascii/hAnsi only; complex-script needs setting too
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), TERM_FONT)


def _set_header_letters(table, header: str) -> None:
    """Replaces the header letters in place, preserving their existing run formatting."""
    for col, letter in enumerate(header):
        cell = table.rows[HEADER_ROW].cells[col]
        runs = [r for p in cell.paragraphs for r in p.runs]
        if runs:
            runs[0].text = letter
            for extra in runs[1:]:
                extra.text = ""
        else:
            _set_term(cell, letter)


def _column_break_paragraph():
    para = OxmlElement("w:p")
    run = OxmlElement("w:r")
    brk = OxmlElement("w:br")
    brk.set(qn("w:type"), "column")
    run.append(brk)
    para.append(run)
    return para


def write_docx_cards(
    cards: list[list[list[str]]],
    template_path: Path,
    output_path: Path,
    free_space: bool,
    free_space_label: str,
    header: str | None = None,
) -> None:
    doc = Document(template_path)
    if not doc.tables:
        raise ValueError(f"No table found in template: {template_path}")

    prototype = copy.deepcopy(doc.tables[0]._tbl)

    # Strip the template's sample content, keeping the section properties (landscape,
    # 1in margins, 2 columns) which live in the trailing sectPr.
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    sect_pr = body.find(qn("w:sectPr"))

    def append(element) -> None:
        if sect_pr is not None:
            sect_pr.addprevious(element)
        else:
            body.append(element)

    for idx, card in enumerate(cards):
        tbl = copy.deepcopy(prototype)
        append(tbl)
        table = doc.tables[-1]

        if header:
            _set_header_letters(table, header)

        for r, row_vals in enumerate(card):
            for c, val in enumerate(row_vals):
                if free_space and r == FREE_ROW and c == FREE_COL:
                    continue  # keep the template's red FREE SPACE cell untouched
                cell = table.rows[r + GRID_ROW_OFFSET].cells[c]
                if r == FREE_ROW and c == FREE_COL:
                    _remove_shading(cell)  # free space disabled: plain term cell
                _set_term(cell, val)

        # Column break sends the next card to the right column, then to a new page.
        if idx < len(cards) - 1:
            append(_column_break_paragraph())
        else:
            append(OxmlElement("w:p"))  # trailing paragraph Word expects after a table

    doc.save(output_path)
